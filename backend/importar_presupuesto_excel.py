"""
Carga masiva del presupuesto de un año (por defecto 2026) por área, desde un Excel
con una hoja por área (~300 filas cada una: categoría, nombre de recurso, detalle,
precio y opcionalmente código contable).

Comportamiento clave (ver plan de diseño en
C:\\Users\\Juan Diego Luque L\\.claude\\plans\\consulta-si-es-posible-jiggly-boole.md):
  - Si un recurso con el mismo nombre ya existe en el catálogo (pre_recurso), se
    reutiliza; nunca se crea un duplicado.
  - Un recurso sin código contable en la planilla se importa igual, sin código
    (se puede completar después desde la UI de GO-Compras).
  - Si la planilla trae un código contable que no existe en el Manual de Cuentas
    (pre_cuenta_matriz_reglas), se descarta el código (no se corta el proceso) y
    se avisa por consola.
  - Por cada área se crea una SolicitudPresupuesto nueva en estado "Aceptado",
    enlazada al PresupuestoAnual del año indicado (se crea si no existe), y cada
    ítem importado queda con estado_aprobacion="Aprobado" para que aparezca de
    inmediato en Programar Compras / Historial.

USO:
    cd backend
    pip install openpyxl        # una sola vez, si no está instalado
    python importar_presupuesto_excel.py --dry-run            # valida sin escribir nada
    python importar_presupuesto_excel.py                      # corre en serio
    python importar_presupuesto_excel.py --solo-hoja "Finanzas"  # una sola área

ANTES DE CORRER EN SERIO: completar la sección de CONFIGURACIÓN más abajo
(ruta del Excel, ID_USER_RESPONSABLE, ID_COLEGIO, nombres de columnas reales,
y los mapeos de hoja→subárea / destino de gasto / subvención).
"""
import argparse
import sys
from datetime import date
from decimal import Decimal, InvalidOperation

try:
    from openpyxl import load_workbook
except ImportError:
    print("Falta la librería 'openpyxl'. Instálala con:  pip install openpyxl")
    sys.exit(1)

from sqlalchemy.orm import Session

from app.db.session import SessionLocal
from app.models import (
    Subarea, SolicitudPresupuesto, PresupuestoDetalle, PresupuestoAnual,
    Recurso, CategoriaRecurso, Subvencion, CuentaMatrizReglas,
)

# ─────────────────────────────────────────────────────────────────────────────
# CONFIGURACIÓN — AJUSTAR ANTES DE CORRER
# ─────────────────────────────────────────────────────────────────────────────

EXCEL_PATH = "presupuesto_2026.xlsx"
ANIO_PRESUPUESTO = 2026

# Se antepone al comentario de cada solicitud creada, para poder detectar en una
# segunda corrida que esa área ya fue importada y no duplicarla.
MARCA_IMPORTACION = "[IMPORT-2026]"

# Usuario que queda como autor de las solicitudes creadas (debe existir en
# auth_usuario) y colegio destino. Completar con los id reales antes de correr.
ID_USER_RESPONSABLE = None   # ej: 1
ID_COLEGIO = None            # ej: 1

# Nombres de columna esperados en cada hoja (ajustar a los encabezados reales del Excel).
COL_CATEGORIA = "Categoría"
COL_RECURSO = "Recurso"
COL_DETALLE = "Detalle"
COL_PRECIO = "Precio"
COL_CANTIDAD = "Cantidad"                  # opcional; si falta o está vacía, se asume 1
COL_FORMATO = "Formato"                    # opcional; si falta, se asume "Unidad"
COL_CODIGO_CONTABLE = "Código Contable"    # opcional
COL_DESTINO = "Destino"                    # opcional; si falta, usa MAPA_DESTINO_POR_HOJA
COL_SUBVENCION = "Subvención"              # opcional; si falta, usa MAPA_SUBVENCION_POR_HOJA

# Nombre de hoja del Excel → nombre EXACTO de la Subárea (org_subarea.nombre).
# Si el nombre de la hoja ya coincide exactamente con la subárea, no hace falta listarla.
MAPA_HOJA_SUBAREA = {
    # "Finanzas": "Finanzas",
}

# Texto libre que puede aparecer en la columna "Destino" → clave que usa
# PresupuestoDetalle.destino_gasto en el sistema.
MAPA_DESTINO_GASTO = {
    "alumnos": "clases(alumno)",
    "sala de clases": "clases(alumno)",
    "estudiantes": "clases(alumno)",
    "administracion": "oficinas(administracion)",
    "administración": "oficinas(administracion)",
    "oficina": "oficinas(administracion)",
    "oficinas": "oficinas(administracion)",
    "premio": "premio/beneficio",
    "beneficio": "premio/beneficio",
    "mantencion": "mantencion/servicio",
    "mantención": "mantencion/servicio",
    "servicio": "mantencion/servicio",
}
DESTINOS_VALIDOS = {"clases(alumno)", "oficinas(administracion)", "premio/beneficio", "mantencion/servicio"}

# Si una hoja completa comparte un mismo destino de gasto (no viene por fila):
MAPA_DESTINO_POR_HOJA = {
    # "Finanzas": "oficinas(administracion)",
}

# Nombre de hoja → nombre_corto de Subvencion (si no viene columna por fila).
MAPA_SUBVENCION_POR_HOJA = {
    # "Finanzas": "SEP",
}

MOTIVO_DEFAULT = "Presupuesto {anio} — importado desde planilla de área."
FECHA_EJECUCION_DEFAULT = date(ANIO_PRESUPUESTO, 1, 1)
TIPO_FECHA_DEFAULT = "anual"

# ─────────────────────────────────────────────────────────────────────────────


class _RecursoPendiente:
    """Marcador liviano usado en --dry-run: representa un recurso que se crearía,
    sin tocar la base de datos."""
    def __init__(self, nombre):
        self.id_recurso = None
        self.nombre = nombre


def cargar_referencias(db: Session):
    subareas = {s.nombre.strip().lower(): s for s in db.query(Subarea).all()}
    categorias = {c.nombre.strip().lower(): c for c in db.query(CategoriaRecurso).filter(CategoriaRecurso.estado == "Activo").all()}
    subvenciones = {s.nombre_corto.strip().lower(): s for s in db.query(Subvencion).all()}
    codigos_validos = {c.codigo for c in db.query(CuentaMatrizReglas).all()}
    cache_recursos = {r.nombre.strip().lower(): r for r in db.query(Recurso).filter(Recurso.estado == "ACTIVO").all()}
    return subareas, categorias, subvenciones, codigos_validos, cache_recursos


def resolver_o_crear_recurso(db, cache_recursos, nombre, descripcion, formato, categoria, dry_run):
    key = nombre.strip().lower()
    existente = cache_recursos.get(key)
    if existente:
        return existente, False

    if dry_run:
        nuevo = _RecursoPendiente(nombre.strip())
        cache_recursos[key] = nuevo
        return nuevo, True

    nuevo = Recurso(
        nombre=nombre.strip(),
        descripcion=descripcion or None,
        formato=formato or None,
        id_cat_recurso=categoria.id_cat_recurso,
        estado="ACTIVO",
    )
    db.add(nuevo)
    db.flush()  # asigna id_recurso sin comprometer la transacción
    cache_recursos[key] = nuevo
    return nuevo, True


def resolver_destino_gasto(nombre_hoja, valor_col):
    if valor_col:
        texto = str(valor_col).strip()
        if texto in DESTINOS_VALIDOS:
            return texto
        clave = texto.lower()
        if clave in MAPA_DESTINO_GASTO:
            return MAPA_DESTINO_GASTO[clave]
    return MAPA_DESTINO_POR_HOJA.get(nombre_hoja)


def parse_decimal(valor):
    if valor is None or valor == "":
        return None
    try:
        return float(Decimal(str(valor).replace(",", ".").strip()))
    except (InvalidOperation, ValueError):
        return None


def procesar_hoja(db: Session, ws, nombre_hoja, referencias, dry_run):
    subareas, categorias, subvenciones, codigos_validos, cache_recursos = referencias

    nombre_subarea = MAPA_HOJA_SUBAREA.get(nombre_hoja, nombre_hoja)
    subarea = subareas.get(nombre_subarea.strip().lower())
    if not subarea:
        print(f"  [OMITIDA] Hoja '{nombre_hoja}': no existe la subárea '{nombre_subarea}'. Revisa MAPA_HOJA_SUBAREA.")
        return

    ya_existe = db.query(SolicitudPresupuesto).filter(
        SolicitudPresupuesto.id_subarea == subarea.id_subarea,
        SolicitudPresupuesto.comentario.like(f"{MARCA_IMPORTACION}%"),
    ).first()
    if ya_existe:
        print(f"  [OMITIDA] Hoja '{nombre_hoja}': ya existe una solicitud importada (id={ya_existe.id_presupuesto}, codigo={ya_existe.codigo}).")
        return

    primera_fila = next(ws.iter_rows(min_row=1, max_row=1))
    col_idx = {celda.value: i for i, celda in enumerate(primera_fila) if celda.value}

    def get(fila, nombre_col):
        idx = col_idx.get(nombre_col)
        if idx is None or idx >= len(fila):
            return None
        return fila[idx].value

    filas_ok = []
    filas_omitidas = []

    for fila in ws.iter_rows(min_row=2):
        if all(c.value is None for c in fila):
            continue
        num_fila = fila[0].row

        nombre_recurso = get(fila, COL_RECURSO)
        if not nombre_recurso or not str(nombre_recurso).strip():
            filas_omitidas.append((num_fila, "sin nombre de recurso"))
            continue

        categoria_nombre = get(fila, COL_CATEGORIA)
        categoria = categorias.get(str(categoria_nombre).strip().lower()) if categoria_nombre else None
        if not categoria:
            filas_omitidas.append((num_fila, f"categoría '{categoria_nombre}' no existe en el catálogo (pre_categoria_recurso)"))
            continue

        precio = parse_decimal(get(fila, COL_PRECIO))
        if precio is None:
            filas_omitidas.append((num_fila, "precio vacío o inválido"))
            continue

        cantidad = parse_decimal(get(fila, COL_CANTIDAD)) or 1

        destino_gasto = resolver_destino_gasto(nombre_hoja, get(fila, COL_DESTINO))
        if not destino_gasto:
            filas_omitidas.append((num_fila, "no se pudo determinar el destino del gasto (revisa MAPA_DESTINO_GASTO / MAPA_DESTINO_POR_HOJA)"))
            continue

        subvencion_raw = get(fila, COL_SUBVENCION)
        nombre_subvencion = str(subvencion_raw).strip().lower() if subvencion_raw else MAPA_SUBVENCION_POR_HOJA.get(nombre_hoja, "").lower()
        subvencion = subvenciones.get(nombre_subvencion) if nombre_subvencion else None

        codigo_raw = get(fila, COL_CODIGO_CONTABLE)
        codigo_cuenta = str(codigo_raw).strip() if codigo_raw not in (None, "") else None
        if codigo_cuenta and codigo_cuenta not in codigos_validos:
            print(f"    [AVISO] fila {num_fila}: código contable '{codigo_cuenta}' no existe en el Manual de Cuentas; se importa sin código.")
            codigo_cuenta = None

        detalle_texto = str(get(fila, COL_DETALLE) or "").strip() or None
        formato_texto = str(get(fila, COL_FORMATO) or "").strip() or "Unidad"

        recurso, _creado = resolver_o_crear_recurso(
            db, cache_recursos,
            nombre=str(nombre_recurso),
            descripcion=detalle_texto,
            formato=formato_texto,
            categoria=categoria,
            dry_run=dry_run,
        )

        filas_ok.append({
            "nombre_producto": str(nombre_recurso).strip(),
            "descripcion": detalle_texto,
            "id_recurso": recurso.id_recurso,
            "codigo_cuenta": codigo_cuenta,
            "formato_unidad": formato_texto,
            "cantidad": cantidad,
            "valor_unitario": precio,
            "valor_unitario_iva": precio,
            "total_iva": round(cantidad * precio, 2),
            "fecha_ejecucion": FECHA_EJECUCION_DEFAULT,
            "tipo_fecha": TIPO_FECHA_DEFAULT,
            "motivo": MOTIVO_DEFAULT.format(anio=ANIO_PRESUPUESTO),
            "destino_gasto": destino_gasto,
            "id_subvencion": subvencion.id_subvencion if subvencion else None,
        })

    recursos_nuevos = sum(1 for f in filas_ok if cache_recursos.get(f["nombre_producto"].strip().lower()) and cache_recursos[f["nombre_producto"].strip().lower()].id_recurso is None)
    print(f"  Hoja '{nombre_hoja}' → subárea '{subarea.nombre}': {len(filas_ok)} filas OK, {len(filas_omitidas)} omitidas.")
    for num_fila, motivo in filas_omitidas:
        print(f"    - fila {num_fila}: {motivo}")

    if not filas_ok:
        return

    if dry_run:
        print(f"    [DRY-RUN] no se escribe nada en la base de datos ({recursos_nuevos} recursos nuevos se crearían).")
        return

    ppto_anual = db.query(PresupuestoAnual).filter(
        PresupuestoAnual.id_colegio == ID_COLEGIO,
        PresupuestoAnual.year == ANIO_PRESUPUESTO,
    ).first()
    if not ppto_anual:
        ppto_anual = PresupuestoAnual(
            id_colegio=ID_COLEGIO,
            year=ANIO_PRESUPUESTO,
            nombre=f"Presupuesto {ANIO_PRESUPUESTO}",
            estado="activo",
            creado_por=ID_USER_RESPONSABLE,
        )
        db.add(ppto_anual)
        db.flush()

    solicitud = SolicitudPresupuesto(
        codigo=f"SP-IMP-{subarea.id_subarea}-{ANIO_PRESUPUESTO}",
        id_user=ID_USER_RESPONSABLE,
        id_subarea=subarea.id_subarea,
        id_colegio=ID_COLEGIO,
        id_presupuesto_anual=ppto_anual.id_presupuesto_anual,
        comentario=f"{MARCA_IMPORTACION} Presupuesto {ANIO_PRESUPUESTO} de {subarea.nombre}, importado desde planilla.",
        estado="Aceptado",
    )
    db.add(solicitud)
    db.flush()

    for det in filas_ok:
        db.add(PresupuestoDetalle(
            id_presupuesto=solicitud.id_presupuesto,
            estado_aprobacion="Aprobado",
            id_subarea=subarea.id_subarea,
            **det,
        ))

    db.commit()
    print(f"    [OK] Solicitud {solicitud.codigo} (id={solicitud.id_presupuesto}) creada con {len(filas_ok)} recursos, todos Aprobados.")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--dry-run", action="store_true", help="Solo valida y muestra el resumen; no escribe en la base de datos.")
    parser.add_argument("--solo-hoja", action="append", default=None, help="Nombre de hoja a procesar (repetible). Si se omite, procesa todas.")
    args = parser.parse_args()

    if not args.dry_run and (ID_USER_RESPONSABLE is None or ID_COLEGIO is None):
        print("ERROR: completa ID_USER_RESPONSABLE e ID_COLEGIO en la sección de CONFIGURACIÓN antes de correr en serio.")
        print("       (para solo validar la planilla sin esos datos, usa --dry-run)")
        sys.exit(1)

    wb = load_workbook(EXCEL_PATH, data_only=True)
    db = SessionLocal()
    try:
        referencias = cargar_referencias(db)
        hojas = args.solo_hoja or wb.sheetnames
        for nombre_hoja in hojas:
            if nombre_hoja not in wb.sheetnames:
                print(f"[AVISO] la hoja '{nombre_hoja}' no existe en el Excel, se omite.")
                continue
            procesar_hoja(db, wb[nombre_hoja], nombre_hoja, referencias, args.dry_run)
        if args.dry_run:
            db.rollback()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


if __name__ == "__main__":
    main()
